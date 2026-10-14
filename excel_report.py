import calendar
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from calendario import CODICI_ASSENZA, durata_turno_minuti, festivita_nazionali, is_festivo


CODICI_TURNO = {"Mattina": "M", "Pomeriggio": "P", "Notte": "N"}
GIORNI_ABBREVIATI = ("L", "M", "M", "G", "V", "S", "D")
TURNI = ("Mattina", "Pomeriggio", "Notte")
PREFISSI_QUOTE = ("M", "P", "N")


def _codice_assenza(tipo: str) -> str:
    try:
        return CODICI_ASSENZA[tipo]
    except KeyError as errore:
        raise ValueError(f"Tipo di assenza non riconosciuto: {tipo}.") from errore


def crea_report_excel(
    struttura: str,
    anno: int,
    mese: int,
    personale: list[dict],
    regole: dict,
    pianificazione: dict,
    assegnazioni: list[dict],
    storico_minuti: dict[str, int] | None = None,
    storico_mesi: dict[int, dict[str, int]] | None = None,
) -> bytes:
    """Crea il calendario mensile modificabile e il riepilogo annuale."""
    workbook = Workbook()
    calendario = workbook.active
    calendario.title = "Calendario"
    riepilogo = workbook.create_sheet("Riepilogo")
    vincoli = workbook.create_sheet("Vincoli")
    info = workbook.create_sheet("Info")
    credito = workbook.create_sheet("Credito")

    personale_per_nome = {
        str(riga.get("Nome", "")).strip(): riga
        for riga in personale
        if str(riga.get("Nome", "")).strip()
    }
    numero_giorni = calendar.monthrange(anno, mese)[1]
    giorni = [date(anno, mese, numero) for numero in range(1, numero_giorni + 1)]
    extra = {
        date.fromisoformat(str(riga["Data"])[:10]): str(riga.get("Descrizione") or "Festività locale")
        for riga in pianificazione.get("festivi_extra", [])
        if riga.get("Data")
    }
    motivi_festivi = festivita_nazionali(anno)
    motivi_festivi.update(extra)

    assenze_per_nome_giorno = {}
    crediti_per_nome_giorno: dict[tuple[str, date], int] = {}
    for assenza in pianificazione.get("assenze", []):
        nome = str(assenza.get("Nome", "")).strip()
        if nome not in personale_per_nome or not assenza.get("Dal") or not assenza.get("Al"):
            continue
        dal = date.fromisoformat(str(assenza["Dal"])[:10])
        al = date.fromisoformat(str(assenza["Al"])[:10])
        tipo = str(assenza.get("Tipo") or "")
        codice = _codice_assenza(tipo)
        turno_previsto = str(assenza.get("Turno previsto") or "")
        credito_giornaliero = (
            durata_turno_minuti(regole, turno_previsto)
            if codice != "R" and turno_previsto
            else 0
        )
        for giorno in giorni:
            if dal <= giorno <= al:
                chiave = (nome, giorno)
                if chiave in assenze_per_nome_giorno:
                    raise ValueError(f"{nome} ha più assenze sovrapposte il {giorno:%d/%m/%Y}.")
                assenze_per_nome_giorno[chiave] = codice
                crediti_per_nome_giorno[chiave] = credito_giornaliero

    turni_per_nome_giorno: dict[tuple[str, date], str] = {}
    for riga in assegnazioni:
        nome = str(riga["Nome"]).strip()
        giorno = date.fromisoformat(str(riga["Data"])[:10])
        codice = CODICI_TURNO.get(riga["Turno"])
        if nome not in personale_per_nome or giorno not in giorni or not codice:
            continue
        chiave = (nome, giorno)
        if chiave in turni_per_nome_giorno:
            raise ValueError(
                f"{nome} ha più di un turno assegnato il {giorno:%d/%m/%Y}."
            )
        if chiave in assenze_per_nome_giorno:
            raise ValueError(
                f"{nome} ha un turno assegnato durante un'assenza il {giorno:%d/%m/%Y}."
            )
        turni_per_nome_giorno[chiave] = codice

    colonne_riepilogo = {}
    colonna_riepilogo = numero_giorni + 2
    intestazioni_riepilogo = [
        "Mattine", "Pomeriggi", "Notti", "Festivi", "Ore lavorate",
        "Ore assenza accreditate", "Ore validate precedenti",
        "Ore progressive", "Obiettivo annuo a oggi", "Scostamento",
    ]
    for offset, intestazione in enumerate(intestazioni_riepilogo):
        colonna = colonna_riepilogo + offset
        colonne_riepilogo[intestazione] = colonna
        calendario.merge_cells(
            start_row=1, start_column=colonna, end_row=2, end_column=colonna
        )
        cella = calendario.cell(row=1, column=colonna, value=intestazione)
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1F4E78")
        cella.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        calendario.column_dimensions[get_column_letter(colonna)].width = (
            20 if intestazione in ("Ore assenza accreditate", "Ore validate precedenti", "Obiettivo annuo a oggi") else 14
        )

    calendario.merge_cells("A1:A2")
    intestazione_nome = calendario["A1"]
    intestazione_nome.value = "Dipendente"
    intestazione_nome.font = Font(bold=True, color="FFFFFF")
    intestazione_nome.fill = PatternFill("solid", fgColor="1F4E78")
    intestazione_nome.alignment = Alignment(horizontal="center", vertical="center")
    calendario.column_dimensions["A"].width = 30

    for indice, giorno in enumerate(giorni, start=2):
        colonna = get_column_letter(indice)
        giorno_settimana = calendario.cell(
            row=1, column=indice, value=GIORNI_ABBREVIATI[giorno.weekday()]
        )
        numero_giorno = calendario.cell(row=2, column=indice, value=giorno.day)
        colore = (
            "1F4E78"
            if giorno.weekday() < 5 and giorno not in motivi_festivi
            else "C65911"
        )
        for cella in (giorno_settimana, numero_giorno):
            cella.font = Font(bold=True, color="FFFFFF")
            cella.fill = PatternFill("solid", fgColor=colore)
            cella.alignment = Alignment(horizontal="center", vertical="center")
        calendario.column_dimensions[colonna].width = 5

    calendario.freeze_panes = "B3"
    calendario.sheet_view.showGridLines = False
    calendario.row_dimensions[1].height = 34
    calendario.row_dimensions[2].height = 24
    validazione = DataValidation(
        type="list", formula1='"M,P,N,R,F,CS"', allow_blank=False
    )
    validazione.error = "Usa M, P, N, R, F oppure CS."
    validazione.errorTitle = "Codice non valido"
    validazione.prompt = "M mattina/malattia · P pomeriggio · N notte · R riposo · F ferie · CS congedo"
    validazione.promptTitle = "Codici calendario"
    calendario.add_data_validation(validazione)

    righe_calendario: dict[str, int] = {}
    righe_credito: dict[str, int] = {}
    credito.append(["Nome", *[giorno.day for giorno in giorni]])
    for indice_persona, (nome, persona) in enumerate(personale_per_nome.items()):
        riga_excel = 3 + indice_persona * 2
        righe_calendario[nome] = riga_excel
        righe_credito[nome] = riga_excel
        calendario.cell(row=riga_excel, column=1, value=nome).alignment = Alignment(
            vertical="center"
        )
        credito.cell(row=riga_excel, column=1, value=nome)
        for indice, giorno in enumerate(giorni, start=2):
            chiave = (nome, giorno)
            codice = assenze_per_nome_giorno.get(
                chiave, turni_per_nome_giorno.get(chiave, "R")
            )
            cella = calendario.cell(row=riga_excel, column=indice, value=codice)
            cella.alignment = Alignment(horizontal="center", vertical="center")
            validazione.add(cella)
            if chiave in assenze_per_nome_giorno:
                colore_codice = {
                    "F": "FCE4D6",
                    "R": "F2F2F2",
                    "M": "F4CCCC",
                    "CS": "D9D2E9",
                }[codice]
            else:
                colore_codice = {
                    "M": "D9EAF7",
                    "P": "E2F0D9",
                    "N": "D9E1F2",
                    "R": "F2F2F2",
                }[codice]
            cella.fill = PatternFill("solid", fgColor=colore_codice)
            cella.font = Font(bold=codice != "R")
            credito.cell(
                row=riga_excel,
                column=indice,
                value=crediti_per_nome_giorno.get(chiave, 0),
            )
        calendario.row_dimensions[riga_excel].height = 22
        if indice_persona < len(personale_per_nome) - 1:
            calendario.row_dimensions[riga_excel + 1].height = 8
            for colonna in range(1, colonna_riepilogo + len(intestazioni_riepilogo)):
                calendario.cell(row=riga_excel + 1, column=colonna).fill = PatternFill(
                    "solid", fgColor="FFFFFF"
                )

        ultima_colonna_giorno = get_column_letter(numero_giorni + 1)
        intervallo = f"$B{riga_excel}:${ultima_colonna_giorno}{riga_excel}"
        intervallo_credito = (
            f"'Credito'!$B{righe_credito[nome]}:"
            f"${ultima_colonna_giorno}{righe_credito[nome]}"
        )
        for codice, titolo in zip(("M", "P", "N"), intestazioni_riepilogo[:3]):
            col_formula = get_column_letter(colonne_riepilogo[titolo])
            if codice == "M":
                calendario[f"{col_formula}{riga_excel}"] = (
                    f'=COUNTIFS({intervallo},"M",{intervallo_credito},0)'
                )
            else:
                calendario[f"{col_formula}{riga_excel}"] = (
                    f'=COUNTIF({intervallo},"{codice}")'
                )

        parti_festivi = []
        for indice, giorno in enumerate(giorni, start=2):
            colonna = get_column_letter(indice)
            festa = is_festivo(giorno, extra)
            notte_festiva = is_festivo(giorno + timedelta(days=1), extra)
            if festa:
                parti_festivi.append(
                    f'COUNTIFS({colonna}{riga_excel},"M",'
                    f"'Credito'!{colonna}{righe_credito[nome]},0)+"
                    f'COUNTIF({colonna}{riga_excel},"P")+'
                    f'COUNTIF({colonna}{riga_excel},"N")'
                )
            elif notte_festiva:
                parti_festivi.append(
                    f'COUNTIF({colonna}{riga_excel},"N")'
                )
        formula_festivi = f"=SUM({','.join(parti_festivi)})" if parti_festivi else "=0"
        calendario.cell(
            row=riga_excel,
            column=colonne_riepilogo["Festivi"],
            value=formula_festivi,
        )
        durata_ore = [
            durata_turno_minuti(regole, turno) / 60 for turno in TURNI
        ]
        ore_lavorate_formula = (
            f"={get_column_letter(colonne_riepilogo['Mattine'])}{riga_excel}*{durata_ore[0]}"
            f"+{get_column_letter(colonne_riepilogo['Pomeriggi'])}{riga_excel}*{durata_ore[1]}"
            f"+{get_column_letter(colonne_riepilogo['Notti'])}{riga_excel}*{durata_ore[2]}"
        )
        calendario.cell(
            row=riga_excel,
            column=colonne_riepilogo["Ore lavorate"],
            value=ore_lavorate_formula,
        )
        formula_credito = (
            f'=SUMPRODUCT((({intervallo}="F")+({intervallo}="M")+'
            f'({intervallo}="CS"))*{intervallo_credito})/60'
        )
        calendario.cell(
            row=riga_excel,
            column=colonne_riepilogo["Ore assenza accreditate"],
            value=formula_credito,
        )
        minuti_precedenti = int((storico_minuti or {}).get(nome, 0))
        calendario.cell(
            row=riga_excel,
            column=colonne_riepilogo["Ore validate precedenti"],
            value=minuti_precedenti / 60,
        )
        col_lavorate = get_column_letter(colonne_riepilogo["Ore lavorate"])
        col_assenze = get_column_letter(colonne_riepilogo["Ore assenza accreditate"])
        col_precedenti = get_column_letter(colonne_riepilogo["Ore validate precedenti"])
        col_progressive = get_column_letter(colonne_riepilogo["Ore progressive"])
        col_obiettivo = get_column_letter(colonne_riepilogo["Obiettivo annuo a oggi"])
        col_scostamento = get_column_letter(colonne_riepilogo["Scostamento"])
        calendario[f"{col_progressive}{riga_excel}"] = (
            f"={col_lavorate}{riga_excel}+{col_assenze}{riga_excel}+{col_precedenti}{riga_excel}"
        )
        if persona.get("Categoria", "Di ruolo") == "Di ruolo":
            ore_settimanali = persona.get("Ore settimanali")
            if ore_settimanali is None:
                ore_settimanali = 38
            obiettivo_ore = (
                float(ore_settimanali)
                * 52
                * giorno_progressivo(anno, mese)
                / (366 if calendar.isleap(anno) else 365)
            )
            calendario[f"{col_obiettivo}{riga_excel}"] = round(obiettivo_ore, 2)
            calendario[f"{col_scostamento}{riga_excel}"] = (
                f"={col_progressive}{riga_excel}-{col_obiettivo}{riga_excel}"
            )
        for colonna in colonne_riepilogo.values():
            calendario.cell(row=riga_excel, column=colonna).number_format = "0.00"

    ultima_riga_personale = 2 if not personale_per_nome else 3 + (len(personale_per_nome) - 1) * 2
    riga_riepilogo_turni = ultima_riga_personale + 2
    for offset, codice in enumerate(("M", "P", "N"), start=0):
        riga = riga_riepilogo_turni + offset
        calendario.cell(row=riga, column=1, value=f"In turno {codice}")
        calendario.cell(row=riga, column=1).font = Font(bold=True)
        for indice in range(2, numero_giorni + 2):
            colonna = get_column_letter(indice)
            riferimenti = [
                (
                    f'COUNTIFS({colonna}{riga_persona},"M",'
                    f"'Credito'!{colonna}{righe_credito[nome]},0)"
                    if codice == "M"
                    else f'COUNTIF({colonna}{riga_persona},"{codice}")'
                )
                for nome, riga_persona in righe_calendario.items()
            ]
            formula = f"=SUM({','.join(riferimenti)})" if riferimenti else "=0"
            calendario.cell(row=riga, column=indice, value=formula)
            calendario.cell(row=riga, column=indice).alignment = Alignment(horizontal="center")
            calendario.cell(row=riga, column=indice).font = Font(bold=True)

    credito.sheet_state = "hidden"
    _crea_riepilogo_annuale(
        riepilogo,
        anno,
        mese,
        personale_per_nome,
        storico_mesi or {},
        storico_minuti or {},
    )
    _crea_vincoli(vincoli, pianificazione)
    _crea_info(info, struttura, anno, mese)
    for sheet in (riepilogo, vincoli):
        for cella in sheet[1]:
            cella.font = Font(bold=True, color="FFFFFF")
            cella.fill = PatternFill("solid", fgColor="1F4E78")
        for colonna in sheet.columns:
            lettera = get_column_letter(colonna[0].column)
            larghezza = max(len(str(cella.value or "")) for cella in colonna)
            sheet.column_dimensions[lettera].width = min(max(larghezza + 2, 12), 30)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def giorno_progressivo(anno: int, mese: int) -> int:
    return date(anno, mese, calendar.monthrange(anno, mese)[1]).timetuple().tm_yday


def _crea_riepilogo_annuale(
    sheet,
    anno: int,
    mese: int,
    personale: dict[str, dict],
    storico_mesi: dict[int, dict[str, int]],
    storico_minuti: dict[str, int],
) -> None:
    mesi = [f"{calendar.month_abbr[indice]}" for indice in range(1, 13)]
    sheet.append(["Dipendente", *mesi, "Ore annue a oggi", "Obiettivo a oggi", "Saldo"])
    giorni_anno = 366 if calendar.isleap(anno) else 365
    progressivo = giorno_progressivo(anno, mese)
    colonna_mese_corrente = get_column_letter(mese + 1)
    colonna_ore_lavorate = get_column_letter(
        calendar.monthrange(anno, mese)[1] + 6
    )
    colonna_credito = get_column_letter(calendar.monthrange(anno, mese)[1] + 7)
    for indice_persona, (nome, persona) in enumerate(personale.items()):
        valori = [nome]
        for mese_numero in range(1, 13):
            if mese_numero < mese:
                minuti = int(storico_mesi.get(mese_numero, {}).get(nome, 0))
                valori.append(minuti / 60 if minuti else None)
            elif mese_numero == mese:
                riga_calendario = 3 + indice_persona * 2
                valori.append(
                    f"=Calendario!{colonna_ore_lavorate}{riga_calendario}"
                    f"+Calendario!{colonna_credito}{riga_calendario}"
                )
            else:
                valori.append(None)
        riga_riepilogo = sheet.max_row + 1
        totale_progressivo_colonna = 14
        obiettivo_colonna = 15
        saldo_colonna = 16
        storico_ore = int(storico_minuti.get(nome, 0)) / 60
        riferimento_mese = f"{colonna_mese_corrente}{riga_riepilogo}"
        sheet.append(valori)
        sheet.cell(
            row=riga_riepilogo,
            column=totale_progressivo_colonna,
            value=f"={riferimento_mese}+{storico_ore}",
        )
        if persona.get("Categoria", "Di ruolo") == "Di ruolo":
            ore_settimanali = persona.get("Ore settimanali")
            if ore_settimanali is None:
                ore_settimanali = 38
            obiettivo = (
                float(ore_settimanali)
                * 52
                * progressivo
                / giorni_anno
            )
            sheet.cell(
                row=riga_riepilogo,
                column=obiettivo_colonna,
                value=round(obiettivo, 2),
            )
            sheet.cell(
                row=riga_riepilogo,
                column=saldo_colonna,
                value=f"=N{riga_riepilogo}-O{riga_riepilogo}",
            )
        else:
            sheet.cell(row=riga_riepilogo, column=obiettivo_colonna, value=None)
            sheet.cell(row=riga_riepilogo, column=saldo_colonna, value=None)
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions


def _crea_vincoli(sheet, pianificazione: dict) -> None:
    sheet.append(["Tipo", "Nome", "Dal/Data", "Al", "Turno previsto / Dettaglio"])
    for assenza in pianificazione.get("assenze", []):
        sheet.append(
            [
                _codice_assenza(str(assenza.get("Tipo") or "")),
                assenza.get("Nome"),
                date.fromisoformat(str(assenza["Dal"])[:10]),
                date.fromisoformat(str(assenza["Al"])[:10]),
                assenza.get("Turno previsto") or "",
            ]
        )
    for turno in pianificazione.get("predeterminati", []):
        giorno = date.fromisoformat(str(turno["Data"])[:10])
        sheet.append(["Predeterminato", turno.get("Nome"), giorno, giorno, turno.get("Turno")])
    for festa in pianificazione.get("festivi_extra", []):
        giorno = date.fromisoformat(str(festa["Data"])[:10])
        sheet.append(["Festività locale", "", giorno, giorno, festa.get("Descrizione")])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=4):
        for cella in row:
            cella.number_format = "DD/MM/YYYY"


def _crea_info(sheet, struttura: str, anno: int, mese: int) -> None:
    sheet.append(["Report turni", f"{struttura} - {mese:02d}/{anno}"])
    sheet.append(["Legenda turni", "M mattina · P pomeriggio · N notte · R riposo"])
    sheet.append(["Legenda assenze", "F ferie · M malattia · CS congedo straordinario · R riposo"])
    sheet.append(["Nota codici", "M mattina è blu; M malattia è rosso."])
    sheet.append(["Riepilogo giornaliero", "Le righe finali del foglio Calendario riportano le coperture M, P e N per giorno."])
    sheet.append(["Riepilogo dipendenti", "Le colonne a destra mostrano turni, festivi, ore lavorate, credito assenze e saldo annuo."])
    sheet.append(["Credito assenze", "Ferie, malattia e congedo accreditano la durata del turno previsto; il riposo (R) vale zero ore."])
    sheet.append(["Obiettivo annuale", "Per il personale di ruolo il riferimento è Ore settimanali × 52; lo scostamento cumulativo è riportato in Riepilogo."])
    sheet.append(["Validazione", "I dati del foglio Excel sono una copia modificabile; non aggiornano l'archivio dell'app."])
    sheet.column_dimensions["A"].width = 25
    sheet.column_dimensions["B"].width = 110
    for cella in sheet[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor="1F4E78")
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
