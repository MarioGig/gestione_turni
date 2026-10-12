import streamlit as st
import pandas as pd
import json
import os
import io
from datetime import date, datetime, timedelta
import calendar
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Configurazione Pagina
st.set_page_config(page_title="Turni Web", page_icon="📅", layout="wide")

st.title("📅 Gestione e Pianificazione Turni Multi-Anno (2026-2050)")
st.caption("Piattaforma Web per la gestione di assenze, turni predefiniti, vincoli, riduzione dinamica e generazione Excel.")

FILE_ASSENZE = "assenze_salvate.json"
FILE_DIPENDENTI = "dipendenti_salvati.json"
FILE_TURNI_FISSI = "turni_fissi_salvati.json"

# 2. Anagrafica Base Default
ANAGRAFICA_DEFAULT = [
    {"nome": "Di Giacinto Gabriella", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Gigante Mario", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 2, "max_festivi_mese": 1, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Ridolfi Alessia", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 1, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": True, "solo_festivi": False},
    {"nome": "Sorgentone Sara", "tipo": "RUOLO", "can_do_night": False, "max_notti_mese": 0, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Fabrizio Alessandra", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Capuani Federica", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Italiani Daniele", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Santilli Federico", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Ferretti Matteo", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Di Resta Alberico", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "Matarazzo Gabriella", "tipo": "RUOLO", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 99, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": False},
    {"nome": "D'Amico Maria Claudia", "tipo": "EXTRA", "can_do_night": False, "max_notti_mese": 0, "max_festivi_mese": 1, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": True},
    {"nome": "Di Giamberardino Giulia", "tipo": "EXTRA", "can_do_night": False, "max_notti_mese": 0, "max_festivi_mese": 1, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": True},
    {"nome": "Di Stefano Alexandra", "tipo": "EXTRA", "can_do_night": False, "max_notti_mese": 0, "max_festivi_mese": 1, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": True},
    {"nome": "Paoletti Giorgio", "tipo": "EXTRA", "can_do_night": False, "max_notti_mese": 0, "max_festivi_mese": 1, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": True},
    {"nome": "Franzone Giuseppina", "tipo": "EXTRA", "can_do_night": True, "max_notti_mese": 99, "max_festivi_mese": 1, "max_turni_mese": 99, "notte_solo_se_festivo": False, "solo_festivi": True}
]

# 3. Funzioni JSON e Calcolo Festività Generiche
def carica_json(filepath, default):
    if os.path.exists(filepath):
        try:
            with open(filepath, "r") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return data
        except Exception:
            pass
    return default

def salva_json(filepath, data):
    try:
        with open(filepath, "w") as f:
            json.dump(data, f, indent=4)
    except Exception:
        pass

def calcola_pasqua(anno):
    """Calcola la domenica di Pasqua usando l'algoritmo di Butcher/Meeus."""
    a = anno % 19
    b = anno // 100
    c = anno % 100
    d = b // 4
    e = b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4
    k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mese = (h + l - 7 * m + 114) // 31
    giorno = ((h + l - 7 * m + 114) % 31) + 1
    return date(anno, mese, giorno)

def get_festivi_anno(anno):
    """Restituisce l'insieme delle festività nazionali italiane fisse e mobili."""
    pasqua = calcola_pasqua(anno)
    pasquetta = pasqua + timedelta(days=1)
    
    festivi = {
        date(anno, 1, 1),   # Capodanno
        date(anno, 1, 6),   # Epifania
        pasqua,             # Pasqua
        pasquetta,          # Pasquetta
        date(anno, 4, 25),  # Liberazione
        date(anno, 5, 1),   # Festa del Lavoro
        date(anno, 6, 2),   # Festa della Repubblica
        date(anno, 8, 15),  # Ferragosto
        date(anno, 11, 1),  # Ognissanti
        date(anno, 12, 8),  # Immacolata
        date(anno, 12, 25), # Natale
        date(anno, 12, 26)  # Santo Stefano
    }
    return festivi

def e_festivo(d):
    festivi_set = get_festivi_anno(d.year)
    return (d.weekday() == 6) or (d in festivi_set)

def get_lettera_giorno(d):
    return ["L", "M", "M", "G", "V", "S", "D"][d.weekday()]

# Inizializzazione Session State
if "dipendenti" not in st.session_state:
    st.session_state["dipendenti"] = carica_json(FILE_DIPENDENTI, ANAGRAFICA_DEFAULT)

if "assenze" not in st.session_state:
    st.session_state["assenze"] = carica_json(FILE_ASSENZE, [
        {"dipendente": "Ridolfi Alessia", "inizio": "2027-06-10", "fine": "2027-06-20", "causale": "F"},
        {"dipendente": "Di Giacinto Gabriella", "inizio": "2027-08-01", "fine": "2027-08-15", "causale": "F"}
    ])

if "turni_fissi" not in st.session_state:
    st.session_state["turni_fissi"] = carica_json(FILE_TURNI_FISSI, [
        {"dipendente": "Gigante Mario", "data": "2027-07-10", "turno": "M"}
    ])

# 4. Algoritmo Turnazione con Priorità alla Mattina e Riduzione Mirata sul Pomeriggio
def esegui_simulazione(dipendenti_list, assenze_list, turni_fissi_list, anno_selezionato):
    elenco = []
    for idx, row in enumerate(dipendenti_list, start=1):
        elenco.append({
            "id": idx, "nome": row["nome"], "tipo": row["tipo"],
            "can_do_night": row["can_do_night"], "max_notti_mese": row["max_notti_mese"],
            "max_festivi_mese": row["max_festivi_mese"], "max_turni_mese": row.get("max_turni_mese", 99),
            "notte_solo_se_festivo": row["notte_solo_se_festivo"],
            "solo_festivi": row["solo_festivi"], "ore_lavorate": 0.0, "festivi_lavorati": 0,
            "ultimo_turno": None, "conteggio_turni": {"M": 0, "P": 0, "N": 0, "R": 0, "F": 0, "MAL": 0},
            "cronologia": {}
        })

    turni_info = {"M": 6.33, "P": 6.33, "N": 12.0, "R": 0.0, "F": 0.0, "MAL": 0.0}

    def turno_valido(dip, t, d_curr):
        last = dip["ultimo_turno"]
        is_f = e_festivo(d_curr)
        if dip["solo_festivi"] and not is_f and t != "R": return False
        if last == "N" and t != "R": return False
        if last == "P" and t == "M": return False
        
        if t in ["M", "P", "N"]:
            tot_turni_mese = sum(1 for d, tur in dip["cronologia"].items() if d.year == d_curr.year and d.month == d_curr.month and tur in ["M", "P", "N"])
            if tot_turni_mese >= dip["max_turni_mese"]:
                return False

        if t == "N":
            if not dip["can_do_night"]: return False
            if dip["notte_solo_se_festivo"] and not is_f: return False
            notti = sum(1 for d, tur in dip["cronologia"].items() if d.year == d_curr.year and d.month == d_curr.month and tur == "N")
            if notti >= dip["max_notti_mese"]: return False

        if is_f and t in ["M", "P", "N"]:
            fest = sum(1 for d, tur in dip["cronologia"].items() if d.year == d_curr.year and d.month == d_curr.month and e_festivo(d) and tur in ["M", "P", "N"])
            if fest >= dip["max_festivi_mese"]: return False
            
        return True

    def seleziona(candidati, codice, num, d_curr):
        validi = [d for d in candidati if turno_valido(d, codice, d_curr)]
        if e_festivo(d_curr) and codice != "R":
            validi.sort(key=lambda x: (0 if x["tipo"] == "EXTRA" else 1, x["festivi_lavorati"], x["conteggio_turni"][codice]))
        else:
            validi.sort(key=lambda x: x["conteggio_turni"][codice])
        return validi[:num]

    def assegna(dip, codice, d_curr):
        dip["ore_lavorate"] += turni_info[codice]
        dip["ultimo_turno"] = codice
        dip["conteggio_turni"][codice] = dip["conteggio_turni"].get(codice, 0) + 1
        dip["cronologia"][d_curr] = codice
        if e_festivo(d_curr) and codice in ["M", "P", "N"]:
            dip["festivi_lavorati"] += 1

    def controlla_assenza(nome, d_curr):
        for a in assenze_list:
            try:
                d_in = datetime.strptime(a["inizio"], "%Y-%m-%d").date()
                d_fi = datetime.strptime(a["fine"], "%Y-%m-%d").date()
                if a["dipendente"] == nome and d_in <= d_curr <= d_fi:
                    return a["causale"]
            except Exception:
                continue
        return None

    def controlla_turno_fisso(nome, d_curr):
        str_d = d_curr.strftime("%Y-%m-%d")
        for tf in turni_fissi_list:
            if tf["dipendente"] == nome and tf["data"] == str_d:
                return tf["turno"]
        return None

    d_curr = date(anno_selezionato, 1, 1)
    d_end = date(anno_selezionato, 12, 31)
    
    while d_curr <= d_end:
        is_f = e_festivo(d_curr)
        
        if is_f:
            n_req, p_req, m_req = 1, 1, 1
        else:
            n_req, p_req, m_req = 1, 2, 6
            
        occupati_temp = set()
        for dip in elenco:
            if controlla_assenza(dip["nome"], d_curr):
                occupati_temp.add(dip["id"])
        num_disponibili = len(elenco) - len(occupati_temp)
        tot_req = n_req + p_req + m_req
        
        if num_disponibili < tot_req:
            deficit = tot_req - num_disponibili
            # Riduciamo PRIMA il pomeriggio
            if not is_f and deficit > 0 and p_req > 0:
                can_reduce_p = min(deficit, p_req)
                p_req -= can_reduce_p
                deficit -= can_reduce_p
            # Se serve, riduciamo la mattina solo dopo aver azzerato i tagli pomeridiani possibili
            if deficit > 0 and m_req > 0:
                can_reduce_m = min(deficit, m_req)
                m_req -= can_reduce_m
                deficit -= can_reduce_m

        fabb = {"N": n_req, "P": p_req, "M": m_req}
        occupati = set()
        
        # 1. Assegnazione Assenze
        for dip in elenco:
            caus = controlla_assenza(dip["nome"], d_curr)
            if caus:
                assegna(dip, caus, d_curr)
                occupati.add(dip["id"])

        # 2. Assegnazione Turni Fissi
        for dip in elenco:
            if dip["id"] not in occupati:
                t_fissato = controlla_turno_fisso(dip["nome"], d_curr)
                if t_fissato:
                    assegna(dip, t_fissato, d_curr)
                    occupati.add(dip["id"])
                    if t_fissato in fabb and fabb[t_fissato] > 0:
                        fabb[t_fissato] -= 1
                
        # 3. Assegnazione prioritaria della MATTINA (M) per prima
        for t_code in ["M", "P", "N"]:
            if fabb[t_code] > 0:
                liberi = [d for d in elenco if d["id"] not in occupati]
                scelti = seleziona(liberi, t_code, fabb[t_code], d_curr)
                for dip in scelti:
                    assegna(dip, t_code, d_curr)
                    occupati.add(dip["id"])
                
        # 4. Tutti i rimanenti vanno in Riposo (R)
        for dip in [d for d in elenco if d["id"] not in occupati]:
            assegna(dip, "R", d_curr)
            
        d_curr += timedelta(days=1)
        
    return elenco

# 5. Generatore Excel Multi-Anno
def genera_excel_bytes(elenco, anno_selezionato, mese_selezionato=0):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    nomi_mesi = ["", "Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno", "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
    font_title = Font(name="Calibri", size=14, bold=True, color="1E293B")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_fest_h = Font(name="Calibri", size=10, bold=True, color="991B1B")
    font_data = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)

    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_festivo = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_assenza = PatternFill(start_color="FEF08A", end_color="FEF08A", fill_type="solid")

    thin = Border(left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0'), top=Side(style='thin', color='E2E8F0'), bottom=Side(style='thin', color='E2E8F0'))
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    mesi_da_generare = range(1, 13) if mese_selezionato == 0 else [mese_selezionato]

    for m in mesi_da_generare:
        ws = wb.create_sheet(title=nomi_mesi[m])
        ws.views.sheetView[0].showGridLines = True
        n_giorni = calendar.monthrange(anno_selezionato, m)[1]
        
        ws.cell(row=1, column=1, value=f"Pianificazione Turni - {nomi_mesi[m]} {anno_selezionato}").font = font_title
        
        # Intestazione Dipendente (righe 3 e 4)
        ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
        c_dip = ws.cell(row=3, column=1, value="Dipendente")
        c_dip.font = font_header; c_dip.fill = fill_header; c_dip.alignment = align_left; c_dip.border = thin
        ws.cell(row=4, column=1).border = thin

        # Intestazioni Giorni (Riga 3: Lettera, Riga 4: Numero)
        for g in range(1, n_giorni + 1):
            d = date(anno_selezionato, m, g)
            is_f = e_festivo(d)
            
            c3 = ws.cell(row=3, column=1 + g, value=get_lettera_giorno(d))
            c3.alignment = align_center; c3.border = thin
            
            c4 = ws.cell(row=4, column=1 + g, value=g)
            c4.alignment = align_center; c4.border = thin
            
            if is_f:
                c3.fill = fill_festivo; c3.font = font_fest_h
                c4.fill = fill_festivo; c4.font = font_fest_h
            else:
                c3.fill = fill_header; c3.font = font_header
                c4.fill = fill_header; c4.font = font_header

        # Intestazioni Totali
        for col_offset, title_text in enumerate(["Ore Mese", "Festivi Mese"], start=n_giorni + 2):
            ws.merge_cells(start_row=3, start_column=col_offset, end_row=4, end_column=col_offset)
            c_tot = ws.cell(row=3, column=col_offset, value=title_text)
            c_tot.font = font_header; c_tot.fill = fill_header; c_tot.alignment = align_center; c_tot.border = thin
            ws.cell(row=4, column=col_offset).border = thin

        # Inserimento Dati Dipendenti con riga vuota di separazione
        r_idx = 5
        for idx, dip in enumerate(elenco):
            ws.cell(row=r_idx, column=1, value=dip["nome"]).font = font_bold
            ws.cell(row=r_idx, column=1).alignment = align_left; ws.cell(row=r_idx, column=1).border = thin
            
            ore_m, fest_m = 0.0, 0
            for g in range(1, n_giorni + 1):
                d = date(anno_selezionato, m, g)
                t = dip["cronologia"][d]
                c = ws.cell(row=r_idx, column=1 + g, value=t)
                c.font = font_data; c.alignment = align_center; c.border = thin
                
                if t in ["F", "MAL"]:
                    c.fill = fill_assenza
                elif e_festivo(d):
                    c.fill = fill_festivo
                    if t in ["M", "P", "N"]: fest_m += 1
                elif idx % 2 == 1:
                    c.fill = fill_zebra
                    
                ore_m += {"M": 6.33, "P": 6.33, "N": 12.0, "R": 0, "F": 0, "MAL": 0}[t]
                
            ws.cell(row=r_idx, column=n_giorni + 2, value=round(ore_m, 2)).font = font_bold
            ws.cell(row=r_idx, column=n_giorni + 2).alignment = align_center; ws.cell(row=r_idx, column=n_giorni + 2).border = thin
            
            ws.cell(row=r_idx, column=n_giorni + 3, value=fest_m).font = font_data
            ws.cell(row=r_idx, column=n_giorni + 3).alignment = align_center; ws.cell(row=r_idx, column=n_giorni + 3).border = thin

            # Riga vuota di separazione
            empty_row = r_idx + 1
            ws.cell(row=empty_row, column=1, value="").border = thin
            for g in range(1, n_giorni + 3):
                ws.cell(row=empty_row, column=1 + g, value="").border = thin

            r_idx += 2

        ws.column_dimensions['A'].width = 24
        for g in range(1, n_giorni + 1):
            ws.column_dimensions[get_column_letter(1 + g)].width = 4.5
        ws.column_dimensions[get_column_letter(n_giorni + 2)].width = 12
        ws.column_dimensions[get_column_letter(n_giorni + 3)].width = 13

    if mese_selezionato == 0:
        ws_sum = wb.create_sheet(title="Riepilogo_Annuale")
        ws_sum.views.sheetView[0].showGridLines = True
        ws_sum.cell(row=1, column=1, value=f"Riepilogo Annuale Turni {anno_selezionato}").font = font_title
        
        headers_sum = ["Dipendente", "Tipo", "Ore Totali", "Mattine (M)", "Pomeriggi (P)", "Notti (N)", "Riposi (R)", "Ferie (F)", "Festivi Lavorati"]
        for col_idx, h in enumerate(headers_sum, 1):
            c = ws_sum.cell(row=3, column=col_idx, value=h)
            c.font = font_header; c.fill = fill_header; c.alignment = align_center; c.border = thin

        for idx, dip in enumerate(elenco):
            r_idx = 4 + idx
            cnt = dip["conteggio_turni"]
            r_vals = [dip["nome"], dip["tipo"], round(dip["ore_lavorate"], 2), cnt.get("M",0), cnt.get("P",0), cnt.get("N",0), cnt.get("R",0), cnt.get("F",0), dip["festivi_lavorati"]]
            for col_idx, val in enumerate(r_vals, 1):
                c = ws_sum.cell(row=r_idx, column=col_idx, value=val)
                c.font = font_bold if col_idx <= 2 else font_data
                c.alignment = align_left if col_idx <= 2 else align_center; c.border = thin
                if idx % 2 == 1: c.fill = fill_zebra

        for col_idx in range(1, len(headers_sum) + 1):
            ws_sum.column_dimensions[get_column_letter(col_idx)].width = 22

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

# 6. Interfaccia Grafica Web (Tabs)
tab1, tab2, tab3, tab4 = st.tabs(["🏖️ Assenze Programmate", "📌 Turni Predefiniti", "👥 Gestione Dipendenti", "🚀 Genera Excel"])

elenco_nomi = [d["nome"] for d in st.session_state["dipendenti"]]

# TAB 1: ASSENZE
with tab1:
    st.subheader("Registrazione Ferie e Permessi")
    if elenco_nomi:
        c1, c2, c3 = st.columns([2, 2, 1])
        with c1:
            dip_sel = st.selectbox("Seleziona Dipendente", elenco_nomi, key="sel_dip_ass")
        with c2:
            d_start = st.date_input("Data Inizio Assenza", date(2027, 6, 10), key="d_start_ass")
            d_end = st.date_input("Data Fine Assenza", date(2027, 6, 20), key="d_end_ass")
        with c3:
            causale_sel = st.selectbox("Causale", ["Ferie (F)", "Malattia (MAL)"], key="caus_ass")
            if st.button("➕ Registra Assenza"):
                st.session_state["assenze"].append({
                    "dipendente": dip_sel,
                    "inizio": d_start.strftime("%Y-%m-%d"),
                    "fine": d_end.strftime("%Y-%m-%d"),
                    "causale": "F" if "Ferie" in causale_sel else "MAL"
                })
                salva_json(FILE_ASSENZE, st.session_state["assenze"])
                st.success("Assenza salvata!")
                st.rerun()

    st.markdown("---")
    st.subheader("Assenze Attualmente Registrate")
    if st.session_state["assenze"]:
        st.dataframe(pd.DataFrame(st.session_state["assenze"]), use_container_width=True)
        if st.button("🗑️ Elimina Tutte le Assenze"):
            st.session_state["assenze"] = []
            salva_json(FILE_ASSENZE, [])
            st.rerun()
    else:
        st.info("Nessuna assenza salvata.")

# TAB 2: TURNI PREDEFINITI (FISSI)
with tab2:
    st.subheader("📌 Assegna Turno Predefinito")
    st.caption("Assegna un turno specifico in una data fissa ad un operatore.")
    
    if elenco_nomi:
        col_t1, col_t2, col_t3, col_t4 = st.columns([2, 2, 2, 1])
        with col_t1:
            dip_tf = st.selectbox("Dipendente", elenco_nomi, key="sel_dip_tf")
        with col_t2:
            data_tf = st.date_input("Data del Turno", date(2027, 7, 10), key="d_tf")
        with col_t3:
            turno_tf = st.selectbox("Turno richiesto", ["M - Mattina", "P - Pomeriggio", "N - Notte", "R - Riposo"], key="t_tf")
        with col_t4:
            st.write("")
            st.write("")
            if st.button("📌 Fissa Turno"):
                codice_t = turno_tf.split(" - ")[0]
                st.session_state["turni_fissi"].append({
                    "dipendente": dip_tf,
                    "data": data_tf.strftime("%Y-%m-%d"),
                    "turno": codice_t
                })
                salva_json(FILE_TURNI_FISSI, st.session_state["turni_fissi"])
                st.success(f"Turno {codice_t} fissato per {dip_tf} il {data_tf}!")
                st.rerun()

    st.markdown("---")
    st.subheader("Elenco Turni Predefiniti Attivi")
    if st.session_state["turni_fissi"]:
        for idx_tf, tf in enumerate(st.session_state["turni_fissi"]):
            c_d, c_dt, c_tr, c_del = st.columns([3, 2, 2, 1])
            c_d.write(f"**{tf['dipendente']}**")
            c_dt.write(f"📅 {tf['data']}")
            c_tr.write(f"Turno: **{tf['turno']}**")
            if c_del.button("❌", key=f"del_tf_{idx_tf}"):
                st.session_state["turni_fissi"].pop(idx_tf)
                salva_json(FILE_TURNI_FISSI, st.session_state["turni_fissi"])
                st.rerun()
        
        if st.button("🗑️ Cancella Tutti i Turni Predefiniti"):
            st.session_state["turni_fissi"] = []
            salva_json(FILE_TURNI_FISSI, [])
            st.rerun()
    else:
        st.info("Nessun turno predefinito configurato.")

# TAB 3: GESTIONE DIPENDENTI E MODIFICA DISPONIBILITÀ
with tab3:
    st.subheader("✏️ Modifica Disponibilità e Vincoli Dipendente Esistente")
    if elenco_nomi:
        dip_mod_sel = st.selectbox("Seleziona il Dipendente da Modificare:", elenco_nomi, key="sel_mod_dip")
        dip_obj = next((d for d in st.session_state["dipendenti"] if d["nome"] == dip_mod_sel), None)
        
        if dip_obj:
            with st.form("form_edit_dipendente"):
                c_m1, c_m2, c_m3 = st.columns([2, 1, 1])
                with c_m1:
                    edit_nome = st.text_input("Nome e Cognome", value=dip_obj["nome"])
                with c_m2:
                    edit_tipo = st.selectbox("Tipo Contratto", ["RUOLO", "EXTRA"], index=0 if dip_obj["tipo"] == "RUOLO" else 1)
                with c_m3:
                    edit_can_night = st.checkbox("Può fare notti?", value=dip_obj["can_do_night"])

                c_m4, c_m5, c_m6, c_m7, c_m8 = st.columns(5)
                with c_m4:
                    edit_max_turni = st.number_input("Max Turni Totali / Mese", min_value=0, max_value=99, value=int(dip_obj.get("max_turni_mese", 99)))
                with c_m5:
                    edit_max_notti = st.number_input("Max Notti / Mese", min_value=0, max_value=99, value=int(dip_obj["max_notti_mese"]))
                with c_m6:
                    edit_max_festivi = st.number_input("Max Festivi / Mese", min_value=0, max_value=99, value=int(dip_obj["max_festivi_mese"]))
                with c_m7:
                    edit_notte_fest = st.checkbox("Notti solo nei festivi", value=dip_obj["notte_solo_se_festivo"])
                with c_m8:
                    edit_solo_fest = st.checkbox("Solo festivi (Extra)", value=dip_obj["solo_festivi"])

                submit_edit = st.form_submit_button("💾 Salva Modifiche Disponibilità")
                if submit_edit:
                    dip_obj["nome"] = edit_nome.strip()
                    dip_obj["tipo"] = edit_tipo
                    dip_obj["can_do_night"] = edit_can_night
                    dip_obj["max_turni_mese"] = int(edit_max_turni)
                    dip_obj["max_notti_mese"] = int(edit_max_notti)
                    dip_obj["max_festivi_mese"] = int(edit_max_festivi)
                    dip_obj["notte_solo_se_festivo"] = edit_notte_fest
                    dip_obj["solo_festivi"] = edit_solo_fest
                    
                    salva_json(FILE_DIPENDENTI, st.session_state["dipendenti"])
                    st.success(f"Disponibilità aggiornate per {edit_nome}!")
                    st.rerun()

    st.markdown("---")
    st.subheader("➕ Aggiungi Nuovo Dipendente")
    with st.form("form_nuovo_dipendente"):
        col_a, col_b, col_c = st.columns([2, 1, 1])
        with col_a:
            nuovo_nome = st.text_input("Nome e Cognome Dipendente")
        with col_b:
            nuovo_tipo = st.selectbox("Tipo Contratto", ["RUOLO", "EXTRA"])
        with col_c:
            can_night = st.checkbox("Può fare notti?", value=True)

        col_d, col_e, col_f, col_g, col_h = st.columns(5)
        with col_d:
            max_turni_new = st.number_input("Max Turni Totali / Mese", min_value=0, max_value=99, value=99)
        with col_e:
            max_notti = st.number_input("Max Notti / Mese", min_value=0, max_value=99, value=99)
        with col_f:
            max_festivi = st.number_input("Max Festivi / Mese", min_value=0, max_value=99, value=99)
        with col_g:
            notte_solo_fest = st.checkbox("Notti solo nei festivi", value=False)
        with col_h:
            solo_festivi = st.checkbox("Solo festivi (Extra)", value=False)

        submit_add = st.form_submit_button("➕ Salva Nuovo Dipendente")
        
        if submit_add:
            if nuovo_nome.strip() != "":
                if any(d["nome"].lower() == nuovo_nome.strip().lower() for d in st.session_state["dipendenti"]):
                    st.error("Esiste già un dipendente con questo nome!")
                else:
                    st.session_state["dipendenti"].append({
                        "nome": nuovo_nome.strip(),
                        "tipo": nuovo_tipo,
                        "can_do_night": can_night,
                        "max_turni_mese": int(max_turni_new),
                        "max_notti_mese": int(max_notti),
                        "max_festivi_mese": int(max_festivi),
                        "notte_solo_se_festivo": notte_solo_fest,
                        "solo_festivi": solo_festivi
                    })
                    salva_json(FILE_DIPENDENTI, st.session_state["dipendenti"])
                    st.success(f"Dipendente '{nuovo_nome}' aggiunto correttamente!")
                    st.rerun()

    st.markdown("---")
    st.subheader("👥 Elenco e Rimozione Dipendenti Attuali")
    
    col_top1, col_top2 = st.columns([3, 1])
    with col_top2:
        if st.button("🔄 Ripristina Anagrafica Default"):
            st.session_state["dipendenti"] = list(ANAGRAFICA_DEFAULT)
            salva_json(FILE_DIPENDENTI, st.session_state["dipendenti"])
            st.rerun()

    for idx, dip in enumerate(st.session_state["dipendenti"]):
        c_nome, c_tipo, c_turni, c_notti, c_fest, c_del = st.columns([3, 1, 2, 2, 2, 1])
        c_nome.write(f"**{dip['nome']}**")
        c_tipo.write(f"`{dip['tipo']}`")
        c_turni.write(f"Turni Max/Mese: {dip.get('max_turni_mese', 99)}")
        c_notti.write(f"Notti Max/Mese: {dip['max_notti_mese'] if dip['can_do_night'] else 'No Notti'}")
        c_fest.write(f"Festivi Max/Mese: {dip['max_festivi_mese']}")
        if c_del.button("❌", key=f"del_dip_{idx}"):
            nome_rimosso = st.session_state["dipendenti"].pop(idx)["nome"]
            salva_json(FILE_DIPENDENTI, st.session_state["dipendenti"])
            st.warning(f"Rimossa/o: {nome_rimosso}")
            st.rerun()

# TAB 4: GENERAZIONE EXCEL CON SELEZIONE DI ANNO E MESE
with tab4:
    st.subheader("📊 Selezione dell'Anno, del Periodo e Calcolo Pianificazione")
    
    col_yr1, col_yr2 = st.columns([1, 2])
    with col_yr1:
        anno_scelto = st.selectbox("🗓️ Seleziona l'Anno di Pianificazione:", list(range(2026, 2051)), index=1) # Default 2027
    
    with col_yr2:
        opzione_periodo = st.radio(
            "Scegli il periodo da esportare nell'Excel:",
            [f"Intero Anno {anno_scelto} (12 Mesi + Foglio Riepilogo)", "Un Singolo Mese Specifico"],
            index=0
        )

    mese_scelto_idx = 0
    nomi_mesi_list = ["Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno", "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]
    if "Singolo Mese" in opzione_periodo:
        mese_str = st.selectbox("Seleziona il mese desiderato:", nomi_mesi_list, index=6)
        mese_scelto_idx = nomi_mesi_list.index(mese_str) + 1

    st.markdown("---")
    if st.button("🚀 Calcola e Genera File Excel", type="primary"):
        with st.spinner(f"Elaborazione turni {anno_scelto} in corso..."):
            risultati = esegui_simulazione(st.session_state["dipendenti"], st.session_state["assenze"], st.session_state["turni_fissi"], anno_selezionato=anno_scelto)
            excel_bytes = genera_excel_bytes(risultati, anno_selezionato=anno_scelto, mese_selezionato=mese_scelto_idx)
            st.session_state["file_excel"] = excel_bytes
            nome_file = f"Pianificazione_Turni_{anno_scelto}_{'Anno_Completo' if mese_scelto_idx == 0 else nomi_mesi_list[mese_scelto_idx-1]}.xlsx"
            st.session_state["nome_file_excel"] = nome_file
            st.success(f"Turnistica per l'anno {anno_scelto} elaborata con successo!")

    if "file_excel" in st.session_state:
        st.download_button(
            label=f"📥 Scarica {st.session_state['nome_file_excel']}",
            data=st.session_state["file_excel"],
            file_name=st.session_state["nome_file_excel"],
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
