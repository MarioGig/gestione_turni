import unittest
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from calendario import genera_calendario
from excel_report import crea_report_excel


def regole(mattina_feriale=0, mattina_festiva=0, pomeriggio=0, notte=0, riposo=11):
    return {
        "turni": [
            {"Turno": "Mattina", "Inizio": "08:00", "Fine": "14:20"},
            {"Turno": "Pomeriggio", "Inizio": "14:00", "Fine": "20:20"},
            {"Turno": "Notte", "Inizio": "20:00", "Fine": "08:00"},
        ],
        "copertura": [
            {"Turno": "Mattina", "Feriale": mattina_feriale, "Festivo": mattina_festiva},
            {"Turno": "Pomeriggio", "Feriale": pomeriggio, "Festivo": pomeriggio},
            {"Turno": "Notte", "Feriale": notte, "Festivo": notte},
        ],
        "riposo_minimo_ore": riposo,
    }


def persona(nome, categoria="Di ruolo"):
    return {
        "Nome": nome,
        "Categoria": categoria,
        "Mattina": True,
        "Pomeriggio": True,
        "Notte": True,
        "Ore settimanali": 38,
    }


class GenerazioneCalendarioTests(unittest.TestCase):
    def test_personale_solo_festivi_non_viene_assegnato_nei_feriali(self):
        personale = [
            persona("Medico"),
            persona("Disponibile festivi", "Solo festivi"),
        ]
        assegnazioni, errori = genera_calendario(
            personale, regole(mattina_feriale=0, mattina_festiva=1), 2025, 1, {}
        )

        self.assertEqual(errori, [])
        self.assertTrue(assegnazioni)
        for assegnazione in assegnazioni:
            self.assertIn(
                assegnazione["Nome"],
                ("Medico", "Disponibile festivi"),
            )
            giorno = date.fromisoformat(assegnazione["Data"])
            if assegnazione["Nome"] == "Disponibile festivi":
                self.assertTrue(giorno.weekday() == 6 or giorno in {
                    date(2025, 1, 1), date(2025, 1, 6)
                })

    def test_assenza_e_turno_predeterminato_sono_rispettati(self):
        personale = [persona("Assente"), persona("Presente")]
        pianificazione = {
            "assenze": [
                {
                    "Nome": "Assente",
                    "Dal": "2025-01-01",
                    "Al": "2025-01-31",
                    "Tipo": "Ferie",
                    "Turno previsto": "Mattina",
                }
            ],
            "predeterminati": [
                {"Nome": "Presente", "Data": "2025-01-10", "Turno": "Mattina"}
            ],
            "quote": [
                {"Nome": "Presente", "M min": 31, "M max": 31}
            ],
        }
        assegnazioni, errori = genera_calendario(
            personale, regole(mattina_feriale=1, mattina_festiva=1, riposo=0),
            2025, 1, pianificazione,
        )

        self.assertEqual(errori, [])
        self.assertEqual(len(assegnazioni), 31)
        self.assertEqual({riga["Nome"] for riga in assegnazioni}, {"Presente"})
        self.assertIn(
            {"Nome": "Presente", "Data": "2025-01-10", "Turno": "Mattina"},
            [
                {"Nome": riga["Nome"], "Data": riga["Data"], "Turno": riga["Turno"]}
                for riga in assegnazioni
            ],
        )

    def test_non_assegna_piu_turni_nella_stessa_giornata(self):
        personale = [persona("Medico")]
        pianificazione = {
            "predeterminati": [
                {"Nome": "Medico", "Data": "2025-01-01", "Turno": "Mattina"},
                {"Nome": "Medico", "Data": "2025-01-01", "Turno": "Notte"},
            ]
        }
        assegnazioni, errori = genera_calendario(
            personale,
            regole(mattina_festiva=1, notte=1),
            2025,
            1,
            pianificazione,
        )

        self.assertEqual(assegnazioni, [])
        self.assertTrue(any("un solo turno al giorno" in errore for errore in errori))

    def test_calendario_infeasible_se_riposo_impedisce_la_copertura(self):
        assegnazioni, errori = genera_calendario(
            [persona("Unica persona")],
            regole(mattina_feriale=1, mattina_festiva=1, pomeriggio=1),
            2025,
            1,
            {},
        )

        self.assertEqual(assegnazioni, [])
        self.assertTrue(any("Non esiste un calendario" in errore for errore in errori))

    def test_report_excel_editabile_con_riepilogo_formula(self):
        personale = [persona("Medico")]
        pianificazione = {"quote": [{"Nome": "Medico", "M min": 1}]}
        assegnazioni, errori = genera_calendario(
            personale,
            regole(mattina_feriale=1, mattina_festiva=1),
            2025,
            1,
            pianificazione,
        )
        self.assertEqual(errori, [])

        contenuto = crea_report_excel(
            "Ospedale",
            2025,
            1,
            personale,
            regole(mattina_feriale=1, mattina_festiva=1),
            pianificazione,
            assegnazioni,
        )
        workbook = load_workbook(BytesIO(contenuto), data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            ["Calendario", "Riepilogo", "Vincoli", "Info", "Credito"],
        )
        self.assertEqual(workbook["Calendario"]["A1"].value, "Dipendente")
        self.assertEqual(workbook["Calendario"]["B1"].value, "M")
        self.assertEqual(workbook["Calendario"]["B2"].value, 1)
        self.assertEqual(workbook["Calendario"]["B3"].value, "M")
        self.assertIn("COUNTIF", workbook["Calendario"]["AH3"].value)
        self.assertEqual(workbook["Calendario"]["A3"].value, "Medico")
        self.assertIn("Calendario!AK3", workbook["Riepilogo"]["B2"].value)
        self.assertIn("=B2+", workbook["Riepilogo"]["N2"].value)

    def test_fra_dipendenti_c_e_una_riga_vuota_e_assenza_diventa_f(self):
        personale = [persona("Medico A"), persona("Medico B")]
        pianificazione = {
            "assenze": [
                {
                    "Nome": "Medico A",
                    "Dal": "2025-01-02",
                    "Al": "2025-01-03",
                    "Tipo": "Ferie",
                    "Turno previsto": "Mattina",
                }
            ]
        }
        contenuto = crea_report_excel(
            "Ospedale", 2025, 1, personale, regole(), pianificazione, []
        )
        calendario = load_workbook(BytesIO(contenuto), data_only=False)["Calendario"]

        self.assertEqual(calendario["C3"].value, "(F)")
        self.assertEqual(calendario["D3"].value, "(F)")
        self.assertIsNone(calendario["A4"].value)
        self.assertEqual(calendario["A5"].value, "Medico B")
        self.assertEqual(calendario["B3"].value, "R")

    def test_riposo_non_accredita_ore_e_festivo_precedente_conta_solo_notte(self):
        personale = [persona("Medico")]
        pianificazione = {
            "assenze": [
                {
                    "Nome": "Medico",
                    "Dal": "2025-01-02",
                    "Al": "2025-01-02",
                    "Tipo": "Riposo",
                }
            ]
        }
        report = crea_report_excel(
            "Ospedale", 2025, 1, personale, regole(), pianificazione, []
        )
        workbook = load_workbook(BytesIO(report), data_only=False)
        calendario = workbook["Calendario"]
        credito = workbook["Credito"]

        self.assertEqual(calendario["C3"].value, "(R)")
        self.assertEqual(credito["C3"].value, 0)
        self.assertIn('COUNTIF(E3,"N")', calendario["AJ3"].value)
        self.assertNotIn('COUNTIF(E3,"M")', calendario["AJ3"].value)
        fine_anno = load_workbook(
            BytesIO(
                crea_report_excel(
                    "Ospedale", 2025, 12, personale, regole(), {}, []
                )
            ),
            data_only=False,
        )["Calendario"]
        self.assertIn('COUNTIF(AF3,"N")', fine_anno["AJ3"].value)
        self.assertNotIn('COUNTIF(AF3,"M")', fine_anno["AJ3"].value)

    def test_riepilogo_annuale_include_mesi_approvati_precedenti(self):
        report = crea_report_excel(
            "Ospedale",
            2025,
            2,
            [persona("Medico")],
            regole(),
            {},
            [],
            storico_minuti={"Medico": 3000},
            storico_mesi={1: {"Medico": 3000}},
        )
        riepilogo = load_workbook(BytesIO(report), data_only=False)["Riepilogo"]

        self.assertEqual(riepilogo["B2"].value, 50)
        self.assertIn("Calendario!AH3", riepilogo["C2"].value)
